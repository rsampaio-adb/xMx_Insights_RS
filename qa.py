#!/usr/bin/env python3
"""QA check: are new/latest models properly represented in dpaas_uccatalog_prd.gmi_ml.xmx_insights?

    python qa.py

Checks (scoped to the "latest model_version per model_category/region/product_category" series
that gmi_ml.model_manager exposes and xmx_insights_data_vw.sql actually surfaces - the same latest-
version rule the view itself applies):

  1. Model coverage - each latest model_key must appear in gmi_ml.channel_contribution (MMM output),
     gmi_ml.raw_data_xmx_models (spend/raw data), and gmi_ml.variable_lineage (variable dim). Missing
     from channel_contribution/raw_data means that whole fact source is silently absent for the
     model in xmx_insights; missing from variable_lineage means every row for the model gets a NULL
     category/channel.
  2. Variable coverage - every variable channel_contribution uses for a latest model should also
     have a variable_lineage row for that model_key. xmx_insights joins channel_contribution to
     variable_lineage with a LEFT JOIN, so an unmatched variable survives with a NULL category ->
     lands in the dashboard's "Other Marketing - Other" bucket instead of its real category - this
     is flagged as an issue. raw_data_xmx_models variables missing from variable_lineage (joined
     with an INNER JOIN, so unmatched ones are silently dropped from xmx_insights) are reported as
     an informational note only, never an issue: raw_data_xmx_models collects more series than any
     one model actually uses, so a variable's absence from a given model's variable_lineage is
     normal, not a data-quality defect.
  3. Channel population - every variable_lineage row for a latest model must have a non-blank
     `channel`, else it renders as "<Category> - Other" in the dashboard.

Outputs:
  cache/qa_results.json           - feeds the "Model QA" tab in report_template.html / index.html
  xMX Insights QA Report.xlsx     - full detail for a human reviewer (Summary + 4 detail tabs)

Requires: requests, openpyxl. Reuses build.py's .env / token / Databricks REST plumbing.
"""
import datetime
import json
import os

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from build import CATALOG, dbx_query, load_env

HERE = os.path.dirname(os.path.abspath(__file__))
QA_CACHE = os.path.join(HERE, "cache", "qa_results.json")
XLSX_OUT = os.path.join(HERE, "xMX Insights QA Report.xlsx")

LATEST_MODEL_CTE = f"""
latest_model AS (
  SELECT md.model_key, md.model_version, md.model_category, md.region, md.geo,
         md.product_cloud, md.customer_group,
         CASE WHEN md.customer_group = 'BIZPRO' THEN 'Acrobat' ELSE md.product_category END AS product_category
  FROM {CATALOG}.gmi_ml.model_manager md
  INNER JOIN (
    SELECT concat(model_category,region,product_category) AS model_series, max(model_version) AS model_version
    FROM {CATALOG}.gmi_ml.model_manager
    GROUP BY ALL
  ) lv ON concat(md.model_category,md.region,md.product_category) = lv.model_series
      AND md.model_version = lv.model_version
  GROUP BY ALL
)
"""

VARIABLE_DIM_CTE = f"""
variable_dim AS (
  SELECT model_key,
         CASE WHEN raw_variable_name IS NULL THEN variable_name ELSE raw_variable_name END AS variable
  FROM {CATALOG}.gmi_ml.variable_lineage
  GROUP BY ALL
)
"""


def run_checks(host, warehouse_id, token):
    def q(sql):
        return dbx_query(host, warehouse_id, token, sql)

    print("Resolving latest model_key per (model_category, region, product_category) series ...")
    latest_models = q(f"""
        WITH {LATEST_MODEL_CTE}
        SELECT * FROM latest_model ORDER BY model_key
    """)
    print(f"  {len(latest_models)} latest model_keys")

    print("Checking model_key coverage across channel_contribution / raw_data_xmx_models / variable_lineage ...")
    coverage = q(f"""
        WITH {LATEST_MODEL_CTE}
        SELECT lm.model_key,
               CASE WHEN cc.model_key IS NOT NULL THEN 'Y' ELSE 'N' END AS in_channel_contribution,
               CASE WHEN rd.model_key IS NOT NULL THEN 'Y' ELSE 'N' END AS in_raw_data,
               CASE WHEN vl.model_key IS NOT NULL THEN 'Y' ELSE 'N' END AS in_variable_lineage
        FROM latest_model lm
        LEFT JOIN (SELECT DISTINCT model_key FROM {CATALOG}.gmi_ml.channel_contribution) cc ON lm.model_key = cc.model_key
        LEFT JOIN (SELECT DISTINCT model_key FROM {CATALOG}.gmi_ml.raw_data_xmx_models) rd ON lm.model_key = rd.model_key
        LEFT JOIN (SELECT DISTINCT model_key FROM {CATALOG}.gmi_ml.variable_lineage) vl ON lm.model_key = vl.model_key
        ORDER BY lm.model_key
    """)

    print("Checking channel_contribution variables against variable_lineage ...")
    missing_cc_vars = q(f"""
        WITH {LATEST_MODEL_CTE}, {VARIABLE_DIM_CTE}
        SELECT cc.model_key, cc.variable, cc.metric, count(*) AS n_rows, round(sum(cc.value), 2) AS total_value
        FROM {CATALOG}.gmi_ml.channel_contribution cc
        INNER JOIN latest_model lm ON cc.model_key = lm.model_key
        LEFT JOIN variable_dim vd ON cc.variable = vd.variable AND cc.model_key = vd.model_key
        WHERE vd.model_key IS NULL
        GROUP BY ALL
        ORDER BY cc.model_key, abs(total_value) DESC
    """)

    print("Checking raw_data_xmx_models variables against variable_lineage ...")
    missing_rd_vars = q(f"""
        WITH {LATEST_MODEL_CTE}, {VARIABLE_DIM_CTE}
        SELECT rd.model_key, rd.variable_name AS variable, count(*) AS n_rows, round(sum(rd.value), 2) AS total_value
        FROM {CATALOG}.gmi_ml.raw_data_xmx_models rd
        INNER JOIN latest_model lm ON rd.model_key = lm.model_key
        LEFT JOIN variable_dim vd ON rd.variable_name = vd.variable AND rd.model_key = vd.model_key
        WHERE vd.model_key IS NULL
        GROUP BY ALL
        ORDER BY rd.model_key, abs(total_value) DESC
    """)
    for r in missing_rd_vars:
        r["is_spend"] = "spend" in r["variable"].lower()

    print("Checking variable_lineage rows for a populated channel value ...")
    missing_channel = q(f"""
        WITH {LATEST_MODEL_CTE}
        SELECT vl.model_key,
               CASE WHEN vl.raw_variable_name IS NULL THEN vl.variable_name ELSE vl.raw_variable_name END AS variable,
               vl.marketing_flag, vl.demand_category, vl.metric
        FROM {CATALOG}.gmi_ml.variable_lineage vl
        INNER JOIN latest_model lm ON vl.model_key = lm.model_key
        WHERE vl.channel IS NULL OR trim(vl.channel) = ''
        ORDER BY vl.model_key
    """)

    return {
        "latest_models": latest_models,
        "coverage": coverage,
        "missing_cc_vars": missing_cc_vars,
        "missing_rd_vars": missing_rd_vars,
        "missing_channel": missing_channel,
    }


def build_model_status(raw):
    model_meta = {m["model_key"]: m for m in raw["latest_models"]}
    cov = {c["model_key"]: c for c in raw["coverage"]}

    cc_by_model = {}
    for r in raw["missing_cc_vars"]:
        if r["metric"] == "arpu":
            continue  # not one of the 3 metrics the dashboard consumes
        d = cc_by_model.setdefault(r["model_key"], {"n_vars": set(), "value": 0.0})
        d["n_vars"].add(r["variable"])
        d["value"] += float(r["total_value"] or 0)

    rd_by_model = {}
    for r in raw["missing_rd_vars"]:
        rd_by_model[r["model_key"]] = rd_by_model.get(r["model_key"], 0) + 1

    chan_by_model = {}
    for r in raw["missing_channel"]:
        chan_by_model[r["model_key"]] = chan_by_model.get(r["model_key"], 0) + 1

    statuses = []
    for mk in sorted(model_meta):
        m = model_meta[mk]
        c = cov.get(mk, {})
        issues = []  # critical: actually affects one of the 6 dashboard views
        notes = []   # informational: doesn't affect the dashboard (metric type it doesn't consume)

        if c.get("in_channel_contribution") != "Y":
            issues.append("Not present in channel_contribution - model produces no MMM "
                           "contribution output in xmx_insights (view will show nothing for it).")
        if c.get("in_raw_data") != "Y":
            issues.append("Not present in raw_data_xmx_models - no spend data joins in; "
                           "Marketing Spend & RoMS and Channel Spend & Performance views will be "
                           "empty for this model.")
        if c.get("in_variable_lineage") != "Y":
            issues.append("Not present in variable_lineage at all - every row for this model "
                           "gets a NULL category/channel in xmx_insights.")

        cc_issue = cc_by_model.get(mk)
        if cc_issue:
            issues.append(
                f"{len(cc_issue['n_vars'])} channel_contribution variable(s) have no "
                f"variable_lineage entry (~${cc_issue['value']:,.0f} of arr/unit contribution) - "
                f"these rows fall into '5. Other Marketing - Other' instead of their real category."
            )

        n_rd_missing = rd_by_model.get(mk, 0)
        if n_rd_missing:
            notes.append(
                f"{n_rd_missing} raw_data_xmx_models variable(s) (incl. spend) have no "
                f"variable_lineage entry - excluded from xmx_insights. Not flagged as an issue: "
                f"not all raw data collected for a model is actually used by it."
            )

        n_chan = chan_by_model.get(mk, 0)
        if n_chan:
            issues.append(f"{n_chan} variable_lineage row(s) for this model have a blank channel - "
                           f"will render as '<Category> - Other'.")

        statuses.append({
            "model_key": mk,
            "model_category": m.get("model_category"),
            "region": m.get("region"),
            "geo": m.get("geo"),
            "product_cloud": m.get("product_cloud"),
            "customer_group": m.get("customer_group"),
            "product_category": m.get("product_category"),
            "model_version": m.get("model_version"),
            "in_channel_contribution": c.get("in_channel_contribution", "N"),
            "in_raw_data": c.get("in_raw_data", "N"),
            "in_variable_lineage": c.get("in_variable_lineage", "N"),
            "status": "OK" if not issues else "ISSUES",
            "issues": issues,
            "notes": notes,
        })
    return statuses


def write_json(statuses, raw):
    n_ok = sum(1 for s in statuses if s["status"] == "OK")
    out = {
        "generated_at": datetime.datetime.now().isoformat(timespec="seconds"),
        "n_models": len(statuses),
        "n_ok": n_ok,
        "n_issues": len(statuses) - n_ok,
        "models": statuses,
    }
    os.makedirs(os.path.dirname(QA_CACHE), exist_ok=True)
    with open(QA_CACHE, "w", encoding="utf-8") as f:
        json.dump(out, f, indent=2)
    print(f"\nWrote {os.path.relpath(QA_CACHE, HERE)}: {n_ok}/{len(statuses)} models OK")


HEADER_FILL = PatternFill("solid", fgColor="1B1B1B")
HEADER_FONT = Font(color="FFFFFF", bold=True)
OK_FILL = PatternFill("solid", fgColor="DFF5E1")
ISSUE_FILL = PatternFill("solid", fgColor="FCE2E2")


def _autosize(ws):
    for col_cells in ws.columns:
        length = max((len(str(c.value)) if c.value is not None else 0) for c in col_cells)
        ws.column_dimensions[get_column_letter(col_cells[0].column)].width = min(max(length + 2, 10), 80)


def _write_sheet(ws, headers, rows):
    ws.append(headers)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    for row in rows:
        ws.append(row)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    _autosize(ws)


def write_excel(statuses, raw):
    wb = Workbook()

    ws = wb.active
    ws.title = "Summary"
    ws.append(["model_key", "status", "region", "geo", "product_cloud", "product_category",
               "model_version", "issue_count", "issues", "notes"])
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    for s in statuses:
        ws.append([s["model_key"], s["status"], s["region"], s["geo"], s["product_cloud"],
                   s["product_category"], s["model_version"], len(s["issues"]),
                   " | ".join(s["issues"]), " | ".join(s["notes"])])
        fill = OK_FILL if s["status"] == "OK" else ISSUE_FILL
        for cell in ws[ws.max_row]:
            cell.fill = fill
        ws.cell(ws.max_row, 9).alignment = Alignment(wrap_text=True, vertical="top")
        ws.cell(ws.max_row, 10).alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for col, width in zip("ABCDEFGH", [34, 9, 8, 8, 13, 16, 13, 11]):
        ws.column_dimensions[col].width = width
    ws.column_dimensions["I"].width = 90
    ws.column_dimensions["J"].width = 60

    ws2 = wb.create_sheet("Model Coverage")
    _write_sheet(ws2, ["model_key", "in_channel_contribution", "in_raw_data", "in_variable_lineage"],
                 [[c["model_key"], c["in_channel_contribution"], c["in_raw_data"], c["in_variable_lineage"]]
                  for c in raw["coverage"]])

    ws3 = wb.create_sheet("Missing Vars - Contribution")
    _write_sheet(ws3, ["model_key", "variable", "metric", "n_rows", "total_value"],
                 [[r["model_key"], r["variable"], r["metric"], r["n_rows"], r["total_value"]]
                  for r in raw["missing_cc_vars"]])

    ws4 = wb.create_sheet("Missing Vars - Raw Data")
    _write_sheet(ws4, ["model_key", "variable", "is_spend", "n_rows", "total_value"],
                 [[r["model_key"], r["variable"], "Y" if r["is_spend"] else "N", r["n_rows"], r["total_value"]]
                  for r in raw["missing_rd_vars"]])

    ws5 = wb.create_sheet("Missing Channel")
    _write_sheet(ws5, ["model_key", "variable", "marketing_flag", "demand_category", "metric"],
                 [[r["model_key"], r["variable"], r["marketing_flag"], r["demand_category"], r["metric"]]
                  for r in raw["missing_channel"]])
    if not raw["missing_channel"]:
        ws5.append(["(no rows - every variable_lineage entry for a latest model has a channel populated)"])

    wb.save(XLSX_OUT)
    print(f"Wrote {os.path.basename(XLSX_OUT)}")


def main():
    env = load_env()
    host = env.get("DATABRICKS_HOST")
    warehouse_id = env.get("DATABRICKS_WAREHOUSE_ID")
    token = env.get("DATABRICKS_PAT") or env.get("DATABRICKS_TOKEN")
    missing = [n for n, v in [("DATABRICKS_HOST", host), ("DATABRICKS_WAREHOUSE_ID", warehouse_id),
                               ("DATABRICKS_PAT or DATABRICKS_TOKEN", token)] if not v]
    if missing:
        raise SystemExit(f"Missing Databricks config: {', '.join(missing)}.")
    host = host.rstrip("/")

    raw = run_checks(host, warehouse_id, token)
    statuses = build_model_status(raw)
    write_json(statuses, raw)
    write_excel(statuses, raw)

    n_ok = sum(1 for s in statuses if s["status"] == "OK")
    print(f"\n{n_ok}/{len(statuses)} latest models clean. Run `python build.py` to refresh the "
          f"dashboard's Model QA tab from cache/qa_results.json.")


if __name__ == "__main__":
    main()
